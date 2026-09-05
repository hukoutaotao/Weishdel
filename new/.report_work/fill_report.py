from __future__ import annotations

import copy
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r"D:\1\new")
SOURCE = ROOT / "自走棋对战系统课程设计报告.docx"
OUTPUT = ROOT / "自走棋对战系统课程设计报告_已填写.docx"


PARAGRAPHS = {
    2: "C++17 + SFML 2.6.2 + Spine Runtime 3.8",
    3: "2026 年 9 月 4 日",
    6: (
        "本项目面向《计算机程序设计基础2》课程设计要求，使用 C++17、SFML 2.6.2 "
        "和 Spine Runtime 3.8 实现一个具有图形界面的自走棋对战系统。系统包含主菜单、帮助、"
        "四张地图选择、三种分队选择、三种 AI 策略、商店刷新与购买、拖拽部署、合成、出售、"
        "复活、提前开战、固定步长战斗、暂停继续、回合结算和结果页面。核心规则集中在 "
        "AutoChessCore 静态库中，SFML 界面只把鼠标操作转换为 GameCommand，AI 也只读取 "
        "ReadOnlyGameView 快照并提交同一种命令。游戏数值由 UTF-8 配置文件读取，角色显示由 "
        "Spine 的 atlas、png 和 skel 资源驱动。当前 Release x64 工程已生成 AutoChessGame.exe、"
        "AutoChessCore.lib 和 AutoChessSpine.lib，并将运行所需 DLL 与 data 目录复制到可执行文件旁。"
        "项目运用了抽象类、继承、虚函数多态、运算符重载、模板容器、智能指针和文件读写。"
    ),
    7: "关键词：自走棋；C++17；SFML；Spine；面向对象；策略模式；固定步长；CMake",
    55: (
        "自走棋系统不是只完成一个算法，而是要让地图选择、经济管理、单位部署、自动战斗、AI "
        "决策和图形显示按同一状态机协同工作。本项目以购买、升级、排列和对战为基本流程，"
        "将规则计算与界面绘制分开，使同一条规则不会分别散落在按钮、AI 和战斗动画中。"
        "课程设计的主要训练目标是把已经学过的 C++ 类、继承、虚函数、容器和文件操作，"
        "应用到一个由多个源文件和外部库组成的完整项目中。"
    ),
    56: (
        "项目的个性化设计体现在四点：第一，五类单位、三种分队和四张地图均由文本配置描述；"
        "第二，进攻型、防守型和路线型 AI 通过同一策略接口工作；第三，Spine 动画把准备、移动、"
        "攻击和死亡状态与核心战斗事件对应；第四，程序优先读取 EXE 旁的 data 目录，便于把整个"
        "生成目录复制到其他位置运行。"
    ),
    60: "可编译：使用 CMake Presets、Ninja、MSVC x64 和 C++17，外部依赖全部放在 third_party 中。",
    61: "可运行：程序创建 2560×1440 窗口，并使用 1280×720 逻辑坐标统一布局和缩放。",
    62: "可维护：核心规则不包含 SFML 头文件；界面、动画、AI、经济和战斗模块按目录分开。",
    63: "可移植：构建后自动复制 SFML DLL、MSVC 运行库和完整 data 目录到 AutoChessGame.exe 旁。",
    64: "可诊断：配置或资源缺失时返回明确错误；Spine 资源加载失败时保留静态单位图形作为回退。",
    65: "可扩展：新增地图和数值主要通过 data 文件完成，新增策略通过实现 IAiStrategy 接口完成。",
    73: (
        "工程由 AutoChessCore、AutoChessSpine 和 AutoChessGame 三个 CMake 目标组成。"
        "AutoChessCore 保存配置、模型、经济、战斗、对局和 AI，不依赖 SFML；"
        "AutoChessSpine 把项目内的 Spine 3.8 C++ runtime 与 spine-sfml 适配层编译为静态库，"
        "并链接 SFML graphics、window、system；AutoChessGame 负责窗口、页面、输入、绘制、"
        "动画状态和运行路径，并链接前两个静态库及三个 SFML 导入目标。"
    ),
    79: "GameApp 从 main.cpp 接收已选择的数据目录，加载 ConfigBundle、三套界面字体和动画资源路径。",
    80: "Screen 将鼠标操作转换为 UiAction，HumanController 再把对局操作转换为 GameCommand。",
    81: "AiController 读取 B 方 ReadOnlyGameView，并委托 IAiStrategy 生成至多一条购买或部署命令。",
    82: "Match::submit 校验命令，Match::step 以 1/60 秒步长推进准备倒计时、BattleSimulation 和回合结算。",
    83: "MatchScreen 根据只读快照绘制 HUD、地图和单位，并把战斗动作序号同步为 Spine 动画状态。",
    84: "CMake 在链接后复制 DLL 与 data；程序运行时优先从 EXE 同级目录读取配置、字体、地图和动画文件。",
    87: (
        "核心配置使用 UTF-8 文本。加载顺序为 game.cfg、units.cfg、factions.cfg，再依次读取 "
        "map_01.map 至 map_04.map。解析器检查重复节、重复字段、未知字段、非法整数或比例、"
        "缺失字段、范围越界和跨文件引用。字体位于 data/fonts，五种单位各有准备区和战斗区两套 "
        "Spine 资源，每套由 atlas、png、skel 三个文件组成。"
    ),
    90: (
        "Unit、Screen、IAiStrategy 和 IPlayerController 都是抽象层。IronGuardUnit、DuelistUnit、"
        "RangerUnit、ArcanistUnit、MedicUnit 通过重写 role() 表现运行时多态；MainMenuScreen、"
        "HelpScreen 和 MatchScreen 通过 Screen 接口统一处理事件与绘制；三种 AI 策略通过 "
        "IAiStrategy 统一选择命令。当前代码在删除主动技能模块后没有保留多重继承，"
        "这是与任务书形式要求尚未完全一致的一点，报告中不虚构已经不存在的类。"
    ),
    95: (
        "为简化值对象比较与日志输出，GridPosition 重载 ==、!=、<，BattlePosition、UnitIdentity "
        "和 PlayerState 中使用的标识对象重载 ==，RoundSummary 与 MatchResultSummary 重载流输出 "
        "operator<<。MatchScreen 内部的 AnimationUnitKey 还重载 ==，用于 unordered_map 的动画实例索引。"
    ),
    99: (
        "BattleSimulation 每帧先检查超时，再清理无效目标；没有目标的单位沿路线移动，随后统一处理"
        "守卫伤害、重新索敌和普通行动。攻击或治疗先形成同帧意图，再集中生效，避免遍历顺序改变结果。"
        "铁卫、决斗者、游侠和奥术师执行物理或法术攻击，医师的普通行为是治疗受伤友军。"
        "界面根据单位移动、攻击和死亡状态播放对应 Spine 动作，但动画只表现结果，不反向修改核心规则。"
    ),
    104: (
        "SFML 创建 2560×1440 物理窗口，并用 1280×720 逻辑视图保持界面比例。BoardTransform 负责"
        "格子和像素坐标换算，MapRenderer 绘制障碍、部署区、守卫和路线，UnitRenderer 提供静态回退图形，"
        "SpineAssetRepository 缓存 atlas、纹理和骨骼数据，UnitAnimationInstance 为每个单位维护独立动作状态。"
        "MatchScreen 只依据 ReadOnlyGameView 更新控件和动画，不直接写入 Match。"
    ),
    108: (
        "// 主循环只收集窗口事件，规则更新固定为每秒 60 次。\n"
        "while (window_.isOpen())\n"
        "{\n"
        "    sf::Event event{};\n"
        "    while (window_.pollEvent(event))\n"
        "    {\n"
        "        handleEvent(event);\n"
        "    }\n\n"
        "    const float frameSeconds = std::min(\n"
        "        frameClock_.restart().asSeconds(), MaximumFrameSeconds);\n"
        "    accumulatorSeconds_ += frameSeconds;\n"
        "    while (accumulatorSeconds_ >= FixedStepSeconds)\n"
        "    {\n"
        "        fixedUpdate();\n"
        "        accumulatorSeconds_ -= FixedStepSeconds;\n"
        "    }\n"
        "    render();\n"
        "}"
    ),
    109: "4.2 抽象类、继承与单位角色",
    110: (
        "// Unit 是抽象基类，五种具体单位通过虚函数报告角色。\n"
        "class Unit\n"
        "{\n"
        "public:\n"
        "    virtual ~Unit() = default;\n"
        "    const UnitDefinition& definition() const noexcept;\n"
        "    virtual UnitRole role() const noexcept = 0;\n"
        "};\n\n"
        "class IronGuardUnit final : public Unit\n"
        "{\n"
        "public:\n"
        "    explicit IronGuardUnit(UnitDefinition definition);\n"
        "    UnitRole role() const noexcept override;\n"
        "};\n\n"
        "// 真人与 AI 控制器都通过同一接口读取快照并返回命令。\n"
        "class IPlayerController\n"
        "{\n"
        "public:\n"
        "    virtual ~IPlayerController() = default;\n"
        "    virtual MapSide side() const noexcept = 0;\n"
        "    virtual std::optional<GameCommand> nextCommand(\n"
        "        const ReadOnlyGameView& view) = 0;\n"
        "};"
    ),
    112: (
        "// variant 限定所有可以提交给 Match 的命令类型。\n"
        "using GameCommandPayload = std::variant<\n"
        "    SelectMapCommand, SelectFactionCommand, SelectAiStrategyCommand,\n"
        "    RefreshShopCommand, PurchaseUnitCommand,\n"
        "    MoveToDeploymentCommand, MoveToReserveCommand,\n"
        "    MergeUnitsCommand, SellUnitCommand, ReviveUnitCommand,\n"
        "    StartCombatCommand>;\n\n"
        "struct GameCommand\n"
        "{\n"
        "    MapSide actor = MapSide::Unknown;\n"
        "    GameCommandPayload payload = RefreshShopCommand{};\n"
        "};"
    ),
    116: "5 系统测试",
    117: "5.1 测试环境与方法",
    119: "5.2 编译与资源复制结果",
    121: "5.3 功能与异常场景验证",
    123: "5.4 测试结论与不足",
    124: (
        "当前副本没有 AutoChessTests 自动测试目标，因此本报告不再沿用旧稿中的 402 项断言数据。"
        "本次以 Release x64 重新构建结果、生成目录文件清单、关键配置检查和现有界面操作截图作为验证依据。"
        "build/msvc-release 中已生成 AutoChessGame.exe、AutoChessCore.lib、AutoChessSpine.lib、"
        "3 个 SFML DLL、3 个 MSVC 运行库和完整 data 目录。data 共包含 3 个 cfg 文件、4 张地图、"
        "1 个中文字体文件以及 30 个 Spine 动画文件。构建与资源复制通过；自动化回归覆盖不足仍需后续补充。"
    ),
    126: (
        "以下截图记录了主菜单、帮助、选择、准备、战斗、暂停和继续等主要交互。截图用于说明界面流程；"
        "当前版本已在同一界面结构上加入 Spine 角色动画，并删除主动技能按钮。"
    ),
    148: "图 6-9 战斗单位、生命状态与目标连线",
    154: "7 系统调试与遇到的困难",
    155: "7.1 从单文件程序到 CMake 多文件项目",
    156: (
        "在此之前我只学习过 C 和 C++，练习通常是一个或少量源文件，没有接触过完整工程。"
        "开始时我不理解头文件、cpp 文件、静态库和可执行文件为什么要分成不同目标，也分不清"
        "“编译错误”和“链接错误”。解决方法是先画出 AutoChessCore、AutoChessSpine、AutoChessGame "
        "三者关系，再在 CMakeLists.txt 中逐个列出源文件、头文件搜索路径和 target_link_libraries。"
        "当某个符号无法找到时，我先确认声明是否可见，再确认实现文件是否加入目标，最后确认目标之间是否链接。"
        "这样把一个模糊的“工程不能运行”拆成配置、编译、链接和运行四个阶段。"
    ),
    157: "7.2 SFML 外部库的编译、链接与运行",
    158: (
        "SFML 是我第一次真正接入的外部库。最容易混淆的是三个层次：include 目录只让编译器找到 .hpp；"
        ".lib 文件让链接器找到函数实现；.dll 文件必须在程序运行时位于 EXE 可搜索的位置。"
        "只完成其中一层时，会分别出现找不到头文件、链接符号失败或双击程序后提示缺少 DLL。"
        "项目最后在 CMake 中把 system、window、graphics 声明为三个 IMPORTED 目标，统一指定 x64 Release "
        "导入库、DLL 和头文件目录，并用 POST_BUILD 命令把三个 DLL 复制到 AutoChessGame.exe 旁。"
        "这一过程让我第一次理解“写 C++ 代码”和“把第三方代码接进项目”是两件不同的工作。"
    ),
    159: "7.3 Spine 动画库与资源文件接入",
    160: (
        "Spine 的接入比 SFML 更复杂，因为它不仅需要头文件和库，还要编译 spine-cpp 的大量源文件，"
        "再与 spine-sfml 适配层和 SFML 链接。开始时我不知道应该把外部源码单独做成库，若直接散放进"
        "游戏目标，很难判断错误来自自己的代码还是运行库。最终建立 AutoChessSpine 静态库，集中设置"
        "include 路径、编译警告和依赖关系。运行时还必须保证每套动画的 atlas、png、skel 版本一致、"
        "文件名一致。SpineAssetRepository 负责加载和缓存，失败时返回错误并使用静态图形回退，"
        "从而避免一个角色资源缺失导致整个程序崩溃。"
    ),
    161: "7.4 路径、字体和动画坐标调试",
    162: "资源路径：程序从其他目录启动时找不到 data。使用 GetModuleFileNameW 获取 EXE 目录，并优先读取同级 data。",
    163: "中文字体：不同电脑不一定有同一种字体。程序依次尝试宋体、华文宋体和随项目附带的 Noto Sans SC。",
    164: "动画比例：Spine 原始骨骼尺寸与棋盘格不一致。按骨骼可见边界计算 scaleForHeight，再针对铁卫死亡动作校准偏移。",
    165: "调试方法：先保存完整错误信息，判断属于配置、编译、链接还是运行；一次只改一个原因，重新构建后再验证。",
    168: "打开 build/msvc-release 目录，确认 AutoChessGame.exe、6 个运行时 DLL 和 data 目录位于同一级。",
    169: "双击 AutoChessGame.exe；资源完整时将出现标题为 AutoChess 的主菜单窗口。",
    170: "首次使用可点击“帮助”，阅读购买、拖拽、合成、出售、复活和开始战斗的方法。",
    171: "点击“开始游戏”，依次选择四张地图之一、玩家分队和 AI 策略。",
    172: "准备阶段点击右侧单位卡购买；按住己方单位拖到蓝色部署格，或拖回备用区和出售区。",
    173: "需要时点击刷新、合成或复活；电脑完成部署后可提前开战，也可等待 45 秒倒计时结束。",
    174: "战斗会自动移动、索敌、攻击或治疗，并播放 relax、move、attack、die 动画；可使用暂停和继续。",
    175: "每回合结算后自动进入下一回合，最多 3 回合；结果页可再来一局或返回主菜单。",
    176: "8.2 构建与运行步骤",
    177: "打开 Visual Studio x64 Developer Command Prompt，并切换到项目根目录。",
    178: "执行 cmake --fresh --preset msvc-release，生成 Ninja Release 构建文件。",
    179: "执行 cmake --build --preset msvc-release --parallel，生成 AutoChessCore、AutoChessSpine 和 AutoChessGame。",
    180: "进入 build/msvc-release，检查 DLL 和 data 已由构建脚本自动复制，然后运行 AutoChessGame.exe。",
    181: "移动程序时应整体复制 AutoChessGame.exe、6 个运行时 DLL 和 data 目录，不能只复制 EXE。",
    183: (
        "我此前只学过 C 和 C++，没有学过 CMake、图形库、骨骼动画库，也没有真正处理过项目之间的依赖。"
        "这次最重要的收获不是又写了几个类，而是第一次看懂一个程序从源文件到可执行文件的完整过程："
        "预处理阶段寻找头文件，编译阶段把 cpp 变成目标文件，链接阶段把自己的目标文件与 .lib 组合，"
        "运行阶段再由系统加载 DLL 和数据资源。理解这四步以后，面对大量报错时不再只是反复修改代码。"
    ),
    184: (
        "另一个收获是学会控制项目边界。核心规则放在不依赖 SFML 的 AutoChessCore 中，"
        "外部 Spine 源码放在 AutoChessSpine 中，界面代码只负责输入与显示；运行资源统一复制到 EXE 旁。"
        "这种组织方式刚开始比写单文件程序麻烦，但后期定位配置、链接、路径和动画问题更清楚。"
        "当前项目仍缺少独立自动测试，也没有保留任务书要求的多重继承，这两点是后续首先需要补齐的内容。"
    ),
    186: (
        "本项目已完成自走棋的地图与分队选择、商店经济、购买与刷新、拖拽部署、合成、出售、复活、"
        "AI 准备、固定步长战斗、回合结算、暂停和结果页面，并成功接入 SFML 与 Spine 3.8。"
        "代码具有抽象类、单继承、多态、运算符重载、文件读写和明确的模块边界，Release x64 构建可以生成"
        "完整运行目录。项目也暴露出两个明确不足：当前副本缺少自动测试目标；删除技能系统后不再有多重继承。"
    ),
    187: (
        "后续应先补充不依赖窗口的核心规则测试，并根据课程要求设计一个具有真实职责的第二接口，"
        "让合适的派生类采用多重继承，而不是为了形式随意增加空基类。之后可继续增加音效、更多地图、"
        "动画混合、动态寻路或网络控制器。扩展时仍应保持 GameCommand 和 ReadOnlyGameView 边界，"
        "避免界面或动画直接修改核心状态。"
    ),
    190: (
        "第 4 章列出了固定步长循环、抽象类、统一命令和运行路径等关键代码并附用途注释。"
        "下表列出当前 src 目录中的 94 个 C++ 源文件与头文件，不包含已经从当前副本删除的测试和技能模块。"
    ),
    194: (
        "以下 4 页保留课程设计任务书与评分表，放在源程序清单之后，便于教师核对课程要求与评分项目。"
    ),
}


TABLES = {
    1: [
        ["功能模块", "需求", "验收边界"],
        ["启动与菜单", "加载配置、字体、SFML DLL 与 Spine 资源；提供开始、帮助和退出。", "关键配置缺失时停止进入主循环，动画失败时可使用静态回退。"],
        ["开局选择", "选择 4 张地图之一、3 个玩家分队之一和 3 种 AI 策略之一。", "选择按状态机顺序进行，非法阶段命令由 Match 拒绝。"],
        ["准备与经济", "收入、败者补助、6 槽商店、刷新、购买、8 槽备用区、部署、合成、出售和复活。", "金币、容量、部署格和单位归属在修改前统一校验。"],
        ["战斗与动画", "单位沿路线移动、索敌、攻击或治疗，到达守卫后结算伤害，并播放 Spine 动作。", "核心按 1/60 秒推进；动画只读取战斗结果。"],
        ["回合与结果", "最多 3 回合，结算死亡、存活、到达守卫和超时状态。", "守卫归零立即结束；最大回合后比较双方守卫。"],
        ["暂停与重开", "暂停冻结核心推进，继续后恢复；结果页支持再来一局和返回菜单。", "暂停期间仍响应窗口事件和界面绘制。"],
        ["AI", "进攻型、防守型、路线型策略完成购买与部署。", "只读取快照并提交统一命令，不直接修改 Match。"],
        ["构建与运行", "CMake 生成三个目标，并复制 DLL 与 data 到 EXE 旁。", "整个生成目录可作为一个完整运行单元。"],
    ],
    2: [
        ["类型", "输入/输出", "具体内容"],
        ["用户输入", "鼠标", "按钮点击、商店购买、拖拽部署或撤回、合成、出售、复活、暂停和继续。"],
        ["配置输入", "UTF-8 文本", "game.cfg、units.cfg、factions.cfg 和 map_01.map 至 map_04.map。"],
        ["资源输入", "字体与动画", "Noto Sans SC 字体，以及每种单位的 atlas、png、skel 文件。"],
        ["程序输出", "SFML 窗口", "菜单、帮助、地图、HUD、商店、单位动画、血条、路线、提示、暂停层和结果页。"],
        ["诊断输出", "标准错误与退出码", "配置路径、资源加载错误和启动失败原因。"],
    ],
    3: [
        ["原则", "实现方式与收益"],
        ["持久单位与战斗单位分离", "OwnedUnit 保存跨回合状态，BattleUnit 只服务当前战斗，便于处理死亡、复活和回合恢复。"],
        ["统一命令入口", "HumanController 与 AiController 都提交 GameCommand，由 Match 负责阶段、归属和资源校验。"],
        ["只读快照", "ReadOnlyGameView 使用值拷贝，不向界面和 AI 暴露 Match 内部可写对象。"],
        ["固定步长", "准备与战斗均按 1/60 秒推进，渲染帧率变化不改变规则顺序。"],
        ["数据驱动", "单位、分队、经济和地图来自配置文件，界面不硬编码规则数值。"],
        ["外部库隔离", "Spine 单独编译为 AutoChessSpine，SFML 通过导入目标链接，依赖关系集中在 CMake 中。"],
    ],
    4: [
        ["文件或目录", "主要内容", "当前规模"],
        ["game.cfg", "回合、准备时间、战斗超时、经济、容量、比例和随机种子", "3 回合；45 秒准备；60 秒战斗；种子 20260814"],
        ["units.cfg", "生命、攻防、射程、速度、价格、守卫伤害和普通行为", "铁卫、决斗者、游侠、奥术师、医师"],
        ["factions.cfg", "守卫、部署上限、价格倍率和单位属性修正", "坚壁、突击、行军 3 个分队"],
        ["map_01～map_02", "双路训练场与多路线峡谷", "11×7 双路；15×9 三路"],
        ["map_03～map_04", "中心对称三路要塞与镜像绕行回廊", "19×11 三路；8×5 绕行路线"],
        ["fonts 与 animations", "中文字体和五种单位的准备/战斗 Spine 资源", "1 个字体；30 个动画文件"],
    ],
    5: [
        ["代表性类", "主要属性（均多于 3 项）", "主要方法（均多于 3 项）"],
        ["GameApp", "dataDirectory_、window_、logicalView_、三套字体、config_、match_、双方控制器、screen_、frameClock_ 等", "run、loadConfiguration、loadFonts、handleEvent、processUiAction、startNewGame、fixedUpdate、render 等"],
        ["Match", "config_、randomEngine_、phase_、currentRound_、地图/分队/AI 选择、双方玩家与商店、battle_ 等", "submit、step、viewFor、selectMap、selectFaction、selectAiStrategy、startCombat、settleCurrentRound 等"],
        ["BattleSimulation", "units_、map_、timeoutFrames_、currentFrame_、finished_、守卫伤害和 summary_", "step、moveUnit、updateTargets、applyBasicActions、applyFrameGuardDamage、finish 等"],
        ["MatchScreen", "只读快照、选择项、商店卡、复活卡、拖拽状态、按钮、HUD、动画仓库和动画实例表", "handleEvent、draw、updateView、updateAnimations、rebuildChoices、finishBasicDrag、drawCombatUnits、syncAnimations 等"],
    ],
    6: [
        ["策略", "决策偏好", "可观察差异"],
        ["进攻型 Offensive", "优先高攻击、攻速与守卫伤害单位，并优先较短路线。", "更快形成进攻压力。"],
        ["防守型 Defensive", "优先生命、防御、法抗和治疗单位，并针对主要威胁路线部署。", "阵容更重视生存与保护。"],
        ["路线型 Route", "综合移动时间、射程、守卫伤害和价格评价单位与路线。", "更重视路线效率。"],
    ],
    7: [
        ["项目", "当前环境或方法"],
        ["操作系统与架构", "Windows x64"],
        ["语言与图形库", "C++17，SFML 2.6.2"],
        ["动画运行库", "Spine Runtime 3.8，项目内编译 spine-cpp 与 spine-sfml"],
        ["编译器", "MSVC 19.51.36256 x64"],
        ["构建工具", "CMake 3.29.2，Ninja 1.12.0，msvc-release 预设"],
        ["构建目标", "AutoChessCore、AutoChessSpine、AutoChessGame"],
        ["验证方式", "Release 构建、输出清单核对、配置检查和主要界面流程截图"],
    ],
    8: [
        ["测试范围", "结果", "依据"],
        ["Release x64 构建", "通过", "生成 AutoChessGame.exe、AutoChessCore.lib、AutoChessSpine.lib"],
        ["运行依赖复制", "通过", "3 个 SFML DLL、3 个 MSVC DLL 与 data 位于 EXE 同级"],
        ["配置文件加载", "通过", "3 个 cfg 文件和 4 张地图均存在于构建输出"],
        ["动画资源完整性", "通过", "5 种单位 × 2 套 × atlas/png/skel，共 30 个文件"],
        ["主要界面流程", "通过", "主菜单、帮助、选择、准备、战斗、暂停与结果截图"],
        ["自动化回归", "未建立", "当前 CMakeLists.txt 中没有测试可执行目标"],
    ],
    9: [
        ["问题现象", "排查方法", "处理结果"],
        ["找不到 SFML 头文件", "检查 include 搜索路径是否指向 third_party/SFML/include", "由导入目标统一提供头文件目录"],
        ["SFML 符号无法链接", "确认 x64 Release 的 graphics/window/system .lib 均已加入", "三个导入目标通过 target_link_libraries 链接"],
        ["双击程序提示缺 DLL", "区分 .lib 与 .dll 的作用并检查 EXE 同级文件", "构建后自动复制 3 个 SFML DLL 和 3 个 MSVC DLL"],
        ["Spine 类型能包含但无实现", "确认 spine-cpp 源文件和 spine-sfml.cpp 是否编入目标", "建立独立 AutoChessSpine 静态库"],
        ["动画文件读取失败", "核对 atlas、png、skel 的目录、文件名和版本", "资源集中放入 data/animations/units"],
        ["换工作目录后找不到 data", "比较当前目录与 EXE 目录", "优先读取 EXE 同级 data"],
        ["中文字体缺失或模糊", "检查系统字体和高分辨率缩放", "多字体回退并采用 1280×720 逻辑视图"],
        ["死亡动画位置偏移", "比较骨骼可见边界、缩放高度和动作锚点", "按动作计算比例并对铁卫死亡动作校准"],
    ],
}


def clone_run_properties(run):
    if run is None or run._element.rPr is None:
        return None
    return copy.deepcopy(run._element.rPr)


def replace_paragraph(paragraph, text: str) -> None:
    sample = next((run for run in paragraph.runs if run.text), None)
    if sample is None and paragraph.runs:
        sample = paragraph.runs[0]
    sample_rpr = clone_run_properties(sample)

    if paragraph.style.name.startswith("Heading"):
        if paragraph.runs:
            paragraph.runs[0].text = text
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(text)
        return

    paragraph.clear()
    run = paragraph.add_run(text)
    if sample_rpr is not None:
        run._element.insert(0, sample_rpr)


def replace_cell(cell, text: str) -> None:
    first = cell.paragraphs[0]
    sample = next((run for run in first.runs if run.text), None)
    if sample is None and first.runs:
        sample = first.runs[0]
    sample_rpr = clone_run_properties(sample)
    first.clear()
    run = first.add_run(text)
    if sample_rpr is not None:
        run._element.insert(0, sample_rpr)
    for paragraph in list(cell.paragraphs[1:]):
        paragraph._element.getparent().remove(paragraph._element)


def fill_table(table, rows: list[list[str]]) -> None:
    while len(table.rows) < len(rows):
        table.add_row()
    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            replace_cell(cell, value)


def source_role(path: str) -> str:
    if path.startswith("src/core/ai/"):
        return "AI 接口、决策工具、策略与控制器"
    if path.startswith("src/core/combat/"):
        return "固定步长战斗、属性结算、索敌与战斗数据"
    if path.startswith("src/core/config/"):
        return "配置解析、类型读取、文件加载与校验"
    if path.startswith("src/core/economy/"):
        return "商店、部署、合成、出售、复活与价格规则"
    if path.startswith("src/core/match/"):
        return "统一命令、对局阶段、回合结算与只读视图"
    if path.startswith("src/core/model/"):
        return "单位、玩家、分队与配置基础模型"
    if path.startswith("src/core/units/"):
        return "单位抽象类、具体单位与角色多态"
    if path.startswith("src/core/map/"):
        return "地图、路线与格子数据"
    if path.startswith("src/core/controllers/"):
        return "玩家控制器抽象接口"
    if path.startswith("src/game/animation/"):
        return "Spine 资源加载、缓存、动作状态与绘制"
    if path.startswith("src/game/controllers/"):
        return "真人输入到统一命令的转换"
    if path.startswith("src/game/rendering/"):
        return "坐标换算、地图与单位静态绘制"
    if path.startswith("src/game/screens/"):
        return "菜单、帮助、对局页面与界面流程"
    if path.startswith("src/game/ui/"):
        return "按钮、主题与界面动作"
    if path in {"src/game/GameApp.cpp", "src/game/GameApp.hpp", "src/game/main.cpp"}:
        return "应用入口、窗口主循环与页面编排"
    if path.startswith("src/game/RuntimePaths"):
        return "可执行文件目录与运行资源路径"
    return "核心公共入口"


def build_source_rows() -> list[list[str]]:
    files = sorted(
        path.relative_to(ROOT).as_posix()
        for path in (ROOT / "src").rglob("*")
        if path.suffix in {".cpp", ".hpp"}
    )
    assert len(files) == 94, len(files)
    rows = [["序号", "文件", "职责"]]
    rows.extend(
        [[str(index), path, source_role(path)] for index, path in enumerate(files, 1)]
    )
    return rows


def main() -> None:
    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)

    for index, text in PARAGRAPHS.items():
        replace_paragraph(document.paragraphs[index], text)

    for index, rows in TABLES.items():
        fill_table(document.tables[index], rows)
    fill_table(document.tables[10], build_source_rows())

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = settings.makeelement(qn("w:updateFields"), {})
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
