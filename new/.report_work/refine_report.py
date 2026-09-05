from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from fill_report import replace_paragraph


ROOT = Path(r"D:\1\new")
REPORT = ROOT / "自走棋对战系统课程设计报告_已填写.docx"
DIAGRAMS = ROOT / ".report_work" / "diagrams"


def replace_inline_image(document, index: int, image_path: Path) -> None:
    inline = document.inline_shapes[index]._inline
    blip = inline.graphic.graphicData.pic.blipFill.blip
    relationship_id = blip.embed
    image_part = document.part.related_parts[relationship_id]
    image_part._blob = image_path.read_bytes()


def remove_numbering(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is not None:
        p_pr.remove(num_pr)
    paragraph.style = "Normal"


def main() -> None:
    document = Document(REPORT)

    replace_paragraph(document.paragraphs[98], "3.5 战斗与结算")
    replace_paragraph(
        document.paragraphs[107],
        "GameApp 把真实帧间隔限制在最大值内，再用累加器按 FixedStepSeconds 重复调用 "
        "fixedUpdate。渲染频率与模拟步长分离后，暂停、动画表现和不同机器帧率不会改变规则顺序。",
    )
    replace_paragraph(
        document.paragraphs[126],
        "以下截图来自功能流程验收阶段，用于展示主菜单、帮助、选择、准备、战斗、暂停和继续。"
        "此后当前代码在相同界面结构上增加了 map_03、map_04 和 Spine 角色动画，并删除主动技能；"
        "因此图 6-3 仍只显示当时的两张地图，图 6-9 至图 6-11 仍保留当时的灰色技能占位区。",
    )

    usage_steps = [
        "打开 build/msvc-release 目录，确认 AutoChessGame.exe、6 个运行时 DLL 和 data 目录位于同一级。",
        "双击 AutoChessGame.exe；资源完整时将出现标题为 AutoChess 的主菜单窗口。",
        "首次使用可点击“帮助”，阅读购买、拖拽、合成、出售、复活和开始战斗的方法。",
        "点击“开始游戏”，依次选择四张地图之一、玩家分队和 AI 策略。",
        "准备阶段点击右侧单位卡购买；按住己方单位拖到蓝色部署格，或拖回备用区和出售区。",
        "需要时点击刷新、合成或复活；电脑完成部署后可提前开战，也可等待 45 秒倒计时结束。",
        "战斗会自动移动、索敌、攻击或治疗，并播放 relax、move、attack、die 动画；可使用暂停和继续。",
        "每回合结算后自动进入下一回合，最多 3 回合；结果页可再来一局或返回主菜单。",
    ]
    for number, (index, text) in enumerate(zip(range(168, 176), usage_steps), 1):
        remove_numbering(document.paragraphs[index])
        replace_paragraph(document.paragraphs[index], f"{number}. {text}")

    build_steps = [
        "打开 Visual Studio x64 Developer Command Prompt，并切换到项目根目录。",
        "执行 cmake --fresh --preset msvc-release，生成 Ninja Release 构建文件。",
        "执行 cmake --build --preset msvc-release --parallel，生成 AutoChessCore、AutoChessSpine 和 AutoChessGame。",
        "进入 build/msvc-release，检查 DLL 和 data 已由构建脚本自动复制，然后运行 AutoChessGame.exe。",
        "移动程序时应整体复制 AutoChessGame.exe、6 个运行时 DLL 和 data 目录，不能只复制 EXE。",
    ]
    for number, (index, text) in enumerate(zip(range(177, 182), build_steps), 1):
        remove_numbering(document.paragraphs[index])
        replace_paragraph(document.paragraphs[index], f"{number}. {text}")

    replace_inline_image(document, 0, DIAGRAMS / "game-flow-current.png")
    replace_inline_image(document, 1, DIAGRAMS / "system-modules-current.png")
    replace_inline_image(document, 2, DIAGRAMS / "class-hierarchy-current.png")

    # The shell contained both an explicit page break and a Heading 1 page break,
    # which produced an otherwise empty page between Chapters 1 and 2.
    page_break_paragraph = document.paragraphs[70]
    page_break_paragraph._element.getparent().remove(page_break_paragraph._element)

    update_fields = document.settings._element.find(qn("w:updateFields"))
    if update_fields is not None:
        update_fields.set(qn("w:val"), "true")
    document.save(REPORT)


if __name__ == "__main__":
    main()
