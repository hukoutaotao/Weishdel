from docx import Document


DOCX_PATH = r"D:\1\new\自走棋对战系统课程设计报告_已填写.docx"


def has_drawing(paragraph):
    return bool(paragraph._p.xpath(".//w:drawing"))


document = Document(DOCX_PATH)

# Figure 6-7 was landing on a page through Word's automatic pagination.  The
# first-line inline picture then rose into the header area.  Make the page
# boundary explicit on the picture paragraph while keeping the caption with it.
caption_index = next(
    index
    for index, paragraph in enumerate(document.paragraphs)
    if "图 6-6" in paragraph.text
)

picture_index = next(
    index
    for index in range(caption_index + 1, len(document.paragraphs))
    if has_drawing(document.paragraphs[index])
)

picture_paragraph = document.paragraphs[picture_index]
picture_paragraph.paragraph_format.page_break_before = True
picture_paragraph.paragraph_format.keep_with_next = True

caption_paragraph = document.paragraphs[picture_index + 1]
caption_paragraph.paragraph_format.keep_with_next = False

document.save(DOCX_PATH)
